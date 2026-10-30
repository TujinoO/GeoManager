import gzip
import io
import json
import tempfile
import zipfile
from dataclasses import replace
from pathlib import Path

import pandas as pd
from django.conf import settings
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group, Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, override_settings

from apps.catalog.analysis import validate_table
from apps.catalog.analysis_visuals import build_visuals
from apps.catalog.analysis_templates import (
    TEMPLATES,
    example_rows,
    genome_example,
    template_bytes,
)
from apps.catalog.models import DataResource
from apps.omics.genome_index import GenomeError, build_index
from apps.omics.models import GenomeIndex


class AnalysisVisualStatisticsTests(SimpleTestCase):
    def test_boxes_keep_zero_exclude_missing_and_use_all_values(self):
        df = pd.DataFrame({"life_form": ["a"] * 5, "height_m": [0, 2, 4, 6, None]})
        visual = build_visuals(df, "traits", "height_m", "life_form")
        self.assertEqual(
            visual["boxes"],
            [
                {
                    "label": "a",
                    "field": "height_m",
                    "n": 4,
                    "min": 0.0,
                    "q1": 1.5,
                    "median": 3.0,
                    "q3": 4.5,
                    "max": 6.0,
                }
            ],
        )

    def test_sample_flow_conserves_count_at_each_stage_including_other(self):
        df = pd.DataFrame(
            {
                "region": [f"r{i}" for i in range(25)],
                "sex": ["unknown"] * 25,
                "material_type": ["leaf"] * 25,
                "is_core": ["false"] * 24 + ["true"],
            }
        )
        visual = build_visuals(df, "samples", None, "sex")
        flow = visual["flows"]
        self.assertEqual(
            sum(x["value"] for x in flow if x["source"].startswith("来源")), 25
        )
        self.assertEqual(
            sum(x["value"] for x in flow if x["source"].startswith("性别")), 25
        )
        self.assertTrue(any(x["source"].endswith("其他") for x in flow))
        self.assertEqual(visual["metrics"][1]["value"], "1")

    def test_molecular_markers_cannot_be_pooled(self):
        df = pd.DataFrame(
            {
                "marker": ["a", "b"],
                "batch": ["b1", "b1"],
                "value": [1, 100],
                "quality_flag": ["pass", "review"],
            }
        )
        self.assertEqual(build_visuals(df, "molecular", "value", "batch")["boxes"], [])
        self.assertEqual(
            build_visuals(df.iloc[:1], "molecular", "value", "batch")["boxes"][0][
                "median"
            ],
            1,
        )

    def test_expression_correlation_paircounts_and_positive_fraction(self):
        df = pd.DataFrame(
            {
                "gene_id": ["a", "b", "c"] * 2,
                "sample_code": ["s1"] * 3 + ["s2"] * 3,
                "group": ["g"] * 6,
                "value": [0, 1, 3, 0, None, 3],
            }
        )
        visual = build_visuals(df, "expression", "value", "group")
        cell = next(
            c
            for c in visual["matrices"][0]["cells"]
            if c["x"] == "s1" and c["y"] == "s2"
        )
        self.assertEqual(cell["n"], 2)
        self.assertIsNone(cell["value"])
        dots = {c["y"]: c for c in visual["dots"]}
        self.assertEqual(dots["a"]["fraction"], 0)
        self.assertEqual(dots["b"]["n"], 1)
        self.assertEqual(dots["c"]["value"], 2)

    def test_function_edges_count_shared_unique_genes_not_annotation_rows(self):
        df = pd.DataFrame(
            {
                "database": ["DB"] * 4,
                "gene_id": ["g1", "g1", "g2", "g1"],
                "term_id": ["t1", "t2", "t2", "t1"],
                "term_name": ["A", "B", "B", "A"],
            }
        )
        visual = build_visuals(df, "function", None, "term_name")
        self.assertEqual(sorted(n["value"] for n in visual["nodes"]), [1, 2])
        self.assertEqual(visual["edges"][0]["value"], 1)
        self.assertEqual(visual["edges"][0]["score"], 0.5)

    def test_variant_bins_keep_chromosomes_separate_and_conserve_counts(self):
        df = pd.DataFrame(
            {
                "chromosome": ["chr01", "chr01", "chr02"],
                "position": [1, 1, 999999],
                "ref": ["A", "A", "AT"],
                "alt": ["G", "AT", "A"],
                "sample_code": ["s1", "s2", "s1"],
                "quality": [10, 20, 30],
            }
        )
        visual = build_visuals(df, "variants", "quality", "chromosome")
        self.assertEqual(sum(b["count"] for b in visual["variantBins"]), 3)
        self.assertEqual(
            sum(
                b["count"] for b in visual["variantBins"] if b["chromosome"] == "chr01"
            ),
            2,
        )
        self.assertEqual(sum(b["count"] for b in visual["variantClasses"]), 3)
        self.assertEqual(len(visual["variantClasses"]), 3)
        self.assertLessEqual(len(visual["variantBins"]), 60)

    def test_visuals_are_independent_of_detail_pagination(self):
        # The endpoint must pass the full filtered dataframe before slicing detail rows.
        from apps.catalog.analysis import analysis_payload
        from types import SimpleNamespace

        df = pd.DataFrame(
            example_rows("traits"), columns=TEMPLATES["traits"]["columns"]
        )
        resource = SimpleNamespace(
            pk=1, name="test", default_visualization={"analysisKind": "traits"}
        )
        a = analysis_payload(resource, {"offset": 0, "limit": 1}, df, len(df))
        b = analysis_payload(resource, {"offset": 30, "limit": 10}, df, len(df))
        self.assertEqual(a["visuals"], b["visuals"])
        self.assertEqual(sum(x["n"] for x in a["visuals"]["boxes"]), 48)


class AnalysisWorkspaceTests(TestCase):
    def setUp(self):
        root = Path(tempfile.mkdtemp(prefix="gpedsp-analysis-test-"))
        self.root = root
        config = replace(
            settings.PROJECT_CONFIG,
            app_data=root / "app",
            research_data_root=root / "research",
        )
        self.override = override_settings(PROJECT_CONFIG=config)
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.owner = get_user_model().objects.create_user(
            username="analysis-owner", password="strong-local-test"
        )
        self.other = get_user_model().objects.create_user(
            username="analysis-other", password="strong-local-test"
        )
        for user in [self.owner, self.other]:
            self.grant(user, "core", "browse_data")
            self.grant(user, "core", "query_data")
        self.grant(self.owner, "catalog", "add_dataresource")
        self.grant(self.owner, "catalog", "export_dataresource")
        self.client.force_login(self.owner)

    def grant(self, user, app, codename):
        user.user_permissions.add(
            Permission.objects.get(content_type__app_label=app, codename=codename)
        )

    def imported(self, kind="traits", rows=None, provenance="synthetic"):
        raw, _, filename = template_bytes(kind, example=True)
        if rows is not None:
            raw = (
                pd.DataFrame(rows, columns=TEMPLATES[kind]["columns"])
                .to_csv(index=False)
                .encode()
            )
        response = self.client.post(
            "/api/analysis/import/",
            {
                "kind": kind,
                "name": "Test " + kind,
                "provenance": provenance,
                "file": SimpleUploadedFile(filename, raw),
            },
        )
        self.assertEqual(response.status_code, 201, response.content)
        return response.json()["resourceId"]

    def post_query(self, id, payload=None):
        return self.client.post(
            f"/api/analysis/resources/{id}/query/",
            data=json.dumps(payload or {}),
            content_type="application/json",
        )

    def genome(self):
        paths = {}
        with zipfile.ZipFile(io.BytesIO(genome_example())) as archive:
            for role, name in [
                ("genome", "genome.fasta"),
                ("gff", "annotation.gff3"),
                ("protein", "protein.faa"),
                ("rna", "rna.fasta"),
                ("cds", "cds.fasta"),
            ]:
                path = self.root / name
                path.write_bytes(archive.read(name))
                paths[role] = path
        return build_index(
            files=paths,
            code="TEST-FG",
            name="Synthetic test",
            user=self.owner,
            provenance="synthetic",
        )

    def test_all_templates_and_examples_import_roundtrip(self):
        for kind in TEMPLATES:
            with self.subTest(kind=kind):
                id = self.imported(kind)
                response = self.post_query(id)
                self.assertEqual(response.status_code, 200, response.content)
                result = response.json()
                self.assertEqual(result["total"], len(example_rows(kind)))
                self.assertEqual(result["provenance"], "synthetic")
                self.assertEqual(result["scope"], "full")

    def test_auto_detects_public_templates_and_preserves_import_kind(self):
        from apps.catalog.analysis_detection import PUBLIC_KINDS

        for kind in PUBLIC_KINDS:
            with self.subTest(kind=kind):
                raw, _, _ = template_bytes(kind, example=True)
                response = self.client.post(
                    "/api/analysis/import/",
                    {
                        "kind": "auto",
                        "name": kind,
                        "provenance": "synthetic",
                        "file": SimpleUploadedFile("unrelated-name.csv", raw),
                    },
                )
                self.assertEqual(response.status_code, 201, response.content)
                self.assertEqual(response.json()["resolvedKind"], kind)
                resource = DataResource.objects.get(pk=response.json()["resourceId"])
                self.assertEqual(resource.default_visualization["analysisKind"], kind)
                self.assertEqual(
                    self.post_query(resource.pk).json()["total"],
                    len(example_rows(kind)),
                )

    def test_source_headers_auto_detect_and_wrong_type_suggests_correction(self):
        cases = [
            ("population", "地点（样方）,种,株数（数量）\nP1,胡杨,4\n"),
            ("community", "Sample,FRic,Shannon\nP1,1.2,0.8\n"),
            ("samples", "DNA样本编号,性别,采集地点\nS1,雌株,塔里木\n"),
            (
                "molecular",
                "result_id,sample_code,marker,unit,value\nR1,S1,DNA,ng/uL,0\n",
            ),
        ]
        before = DataResource.objects.count()
        for expected, text in cases:
            for chosen in ["auto", "individual"]:
                response = self.client.post(
                    "/api/analysis/import/",
                    {
                        "kind": chosen,
                        "name": "source",
                        "provenance": "observed",
                        "validateOnly": "true",
                        "file": SimpleUploadedFile("data.csv", text.encode()),
                    },
                )
                if chosen == "auto":
                    self.assertEqual(response.status_code, 201, response.content)
                    self.assertEqual(response.json()["resolvedKind"], expected)
                else:
                    self.assertEqual(response.status_code, 400)
                    self.assertIn(
                        TEMPLATES[expected]["name"], response.json()["issues"][0]
                    )
        self.assertEqual(before, DataResource.objects.count())

    def test_auto_rejects_unknown_and_ambiguous_headers(self):
        for text in [
            "name,value\na,1\n",
            "collection_code,species,sample_code\nC1,胡杨,S1\n",
        ]:
            response = self.client.post(
                "/api/analysis/import/",
                {
                    "kind": "auto",
                    "name": "unknown",
                    "provenance": "observed",
                    "file": SimpleUploadedFile("个体数据.csv", text.encode()),
                },
            )
            self.assertEqual(response.status_code, 400)
            self.assertIn("手动选择", response.json()["detail"])

    def test_ecology_group_recognizes_subtype_and_rejects_other_domains(self):
        for expected in [
            "individual",
            "population",
            "community",
            "traits",
            "molecular",
        ]:
            raw, _, _ = template_bytes(expected, example=True)
            response = self.client.post(
                "/api/analysis/import/",
                {
                    "kind": "ecology",
                    "name": "ecology",
                    "provenance": "synthetic",
                    "validateOnly": "true",
                    "file": SimpleUploadedFile("table.csv", raw),
                },
            )
            self.assertEqual(
                response.status_code,
                400 if expected == "molecular" else 201,
                response.content,
            )
            if expected != "molecular":
                self.assertEqual(response.json()["resolvedKind"], expected)

    def test_molecular_cannot_be_misclassified_as_samples_even_with_unique_sample(self):
        response = self.client.post(
            "/api/analysis/import/",
            {
                "kind": "samples",
                "name": "wrong",
                "provenance": "observed",
                "file": SimpleUploadedFile(
                    "sample.csv",
                    b"result_id,sample_code,marker,unit,value\nR1,S1,DNA,ng/uL,0\n",
                ),
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("分子", response.json()["issues"][0])

    def test_template_xlsx_dictionary_and_empty_csv(self):
        from openpyxl import load_workbook

        for kind in TEMPLATES:
            raw, _, _ = template_bytes(kind, file_format="xlsx", example=True)
            wb = load_workbook(io.BytesIO(raw), read_only=True)
            self.assertEqual(wb.sheetnames, ["data", "字段说明"])
            self.assertEqual(wb["data"].max_row, len(example_rows(kind)) + 1)
            wb.close()
        response = self.client.get("/api/analysis/templates/traits/file/?format=csv")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.content.decode("utf-8-sig").splitlines()), 1)
        self.assertEqual(
            self.client.get("/api/analysis/templates/not-real/file/").status_code, 400
        )

    def test_every_excel_example_validates_through_upload_endpoint(self):
        before = DataResource.objects.count()
        for kind in TEMPLATES:
            with self.subTest(kind=kind):
                raw, _, name = template_bytes(kind, file_format="xlsx", example=True)
                response = self.client.post(
                    "/api/analysis/import/",
                    {
                        "kind": kind,
                        "name": "Excel " + kind,
                        "provenance": "synthetic",
                        "validateOnly": "true",
                        "file": SimpleUploadedFile(name, raw),
                    },
                )
                self.assertEqual(response.status_code, 201, response.content)
                self.assertTrue(response.json()["valid"])
                self.assertEqual(response.json()["rowCount"], len(example_rows(kind)))
        self.assertEqual(DataResource.objects.count(), before)

    def test_requires_explicit_provenance_and_permission(self):
        raw, _, _ = template_bytes("traits", example=True)
        response = self.client.post(
            "/api/analysis/import/",
            {
                "kind": "traits",
                "name": "test",
                "file": SimpleUploadedFile("test.csv", raw),
            },
        )
        self.assertEqual(response.status_code, 400)
        self.client.force_login(self.other)
        self.assertEqual(self.client.post("/api/analysis/import/", {}).status_code, 403)

    def test_validate_only_does_not_write_resource(self):
        raw, _, _ = template_bytes("traits", example=True)
        before = DataResource.objects.count()
        response = self.client.post(
            "/api/analysis/import/",
            {
                "kind": "traits",
                "name": "validation",
                "provenance": "observed",
                "validateOnly": "true",
                "file": SimpleUploadedFile("test.csv", raw),
            },
        )
        self.assertEqual(response.status_code, 201)
        self.assertIsNone(response.json()["resourceId"])
        self.assertEqual(before, DataResource.objects.count())

    def test_duplicate_keys_and_fraction_range_and_nonfinite_rejected(self):
        rows = example_rows("traits")[:2]
        rows[1][0] = rows[0][0]
        rows[0][6] = 72
        rows[1][5] = "inf"
        df = pd.DataFrame(rows, columns=TEMPLATES["traits"]["columns"])
        issues, _ = validate_table(df, "traits")
        self.assertTrue(any("重复" in i for i in issues))
        self.assertTrue(any("0–1" in i for i in issues))
        self.assertTrue(any("有限数值" in i for i in issues))

    def test_mixed_expression_units_and_assemblies_are_rejected(self):
        rows = example_rows("expression")[:2]
        rows[1][0] = "other"
        rows[1][-1] = "count"
        issues, _ = validate_table(
            pd.DataFrame(rows, columns=TEMPLATES["expression"]["columns"]), "expression"
        )
        self.assertTrue(any("单位" in i for i in issues))
        self.assertTrue(any("组装" in i for i in issues))

    def test_filters_statistics_pagination_and_export_use_same_population(self):
        id = self.imported()
        payload = {
            "filters": [{"field": "region", "op": "eq", "value": "塔里木"}],
            "metric": "height_m",
            "groupBy": "species",
            "offset": 2,
            "limit": 3,
            "xField": "height_m",
            "yField": "leaf_water_fraction",
        }
        response = self.post_query(id, payload)
        self.assertEqual(response.status_code, 200, response.content)
        result = response.json()
        self.assertEqual(result["filteredCount"], 16)
        self.assertEqual(len(result["rows"]), 3)
        self.assertEqual(result["pairedCount"], 16)
        self.assertEqual(sum(b["count"] for b in result["histogram"]), 16)
        export = self.client.post(
            f"/api/analysis/resources/{id}/export/",
            data=json.dumps(payload),
            content_type="application/json",
        )
        self.assertEqual(export.status_code, 200)
        self.assertEqual(len(pd.read_csv(io.BytesIO(export.content))), 16)
        self.assertEqual(result["scope"], "filtered")

    def test_empty_filter_has_no_invented_statistics(self):
        id = self.imported()
        result = self.post_query(id, {"keyword": "不存在的对象"}).json()
        self.assertEqual(result["filteredCount"], 0)
        self.assertEqual(result["rows"], [])
        self.assertEqual(result["histogram"], [])
        self.assertIsNone(result["correlation"])

    def test_zero_is_preserved_in_distribution(self):
        rows = example_rows("traits")[:2]
        rows[0][5] = 0
        id = self.imported(rows=rows)
        result = self.post_query(id, {"metric": "height_m"}).json()
        profile = next(f for f in result["fields"] if f["name"] == "height_m")
        self.assertEqual(profile["min"], 0)
        self.assertEqual(profile["n"], 2)

    def test_query_rejects_unbounded_sql_and_unknown_fields(self):
        id = self.imported()
        for payload in [
            {"limit": 99999},
            {"metric": "bad"},
            {"sql": "DROP TABLE x"},
            {"filters": [{"field": "height_m", "op": "gte", "value": "nan"}]},
            {"filters": [{"field": "not-field", "op": "eq", "value": "x"}]},
        ]:
            with self.subTest(payload=payload):
                self.assertEqual(self.post_query(id, payload).status_code, 400)

    def test_export_escapes_formula_strings(self):
        rows = example_rows("traits")[:1]
        rows[0][2] = "=SUM(1,2)"
        id = self.imported(rows=rows)
        response = self.client.post(
            f"/api/analysis/resources/{id}/export/",
            data="{}",
            content_type="application/json",
        )
        self.assertIn("'=SUM", response.content.decode("utf-8-sig"))

    def test_resource_permissions_apply_to_query_export_and_views(self):
        id = self.imported()
        save = self.client.post(
            "/api/analysis/views/",
            data=json.dumps({"name": "owner-only", "resourceId": id, "query": {}}),
            content_type="application/json",
        )
        self.assertEqual(save.status_code, 201, save.content)
        self.client.force_login(self.other)
        self.assertEqual(self.post_query(id).status_code, 403)
        self.assertEqual(self.client.get("/api/analysis/views/").json()["items"], [])
        response = self.client.post(
            f"/api/analysis/resources/{id}/export/",
            data="{}",
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 403)

    def test_export_requires_separate_permission_even_when_resource_accessible(self):
        id = self.imported()
        group = Group.objects.create(name="analysis-shared")
        DataResource.objects.get(pk=id).access_groups.add(group)
        self.other.groups.add(group)
        self.client.force_login(self.other)
        self.assertEqual(self.post_query(id).status_code, 200)
        self.assertEqual(
            self.client.post(
                f"/api/analysis/resources/{id}/export/",
                data="{}",
                content_type="application/json",
            ).status_code,
            403,
        )

    def test_unauthenticated_endpoints_are_json_401(self):
        self.client.logout()
        for url in [
            "/api/analysis/templates/",
            "/api/analysis/views/",
            "/api/omics/assemblies/",
            "/api/omics/assemblies/1/summary/",
        ]:
            response = self.client.get(url)
            self.assertEqual(response.status_code, 401)
            self.assertIn("detail", response.json())

    def test_genome_full_index_search_alias_and_model(self):
        index = self.genome()
        summary = self.client.get(
            f"/api/omics/assemblies/{index.assembly_id}/summary/"
        ).json()
        self.assertEqual(summary["geneCount"], 60)
        self.assertEqual(summary["proteinCount"], 60)
        self.assertEqual(summary["chromosomeCount"], 3)
        response = self.client.get(
            f"/api/omics/assemblies/{index.assembly_id}/genes/?q=DEMO-P01001"
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["items"][0]["id"], "DEMO-G01001")
        partial = self.client.get(
            f"/api/omics/assemblies/{index.assembly_id}/genes/?q=demo-p01"
        )
        self.assertEqual(partial.status_code, 200)
        self.assertEqual(partial.json()["total"], 20)
        detail = self.client.get(
            f"/api/omics/assemblies/{index.assembly_id}/genes/DEMO-G01001/"
        ).json()
        self.assertEqual(len(detail["sequences"]), 3)
        self.assertEqual(
            {f["type"] for f in detail["features"]}, {"gene", "mRNA", "exon", "CDS"}
        )

    def test_genome_closed_coordinates_and_reverse_complement(self):
        index = self.genome()
        base = f"/api/omics/assemblies/{index.assembly_id}/sequence/"
        plus = self.client.get(
            base, {"seqid": "chr01", "start": 1, "end": 3, "strand": "+"}
        ).json()
        minus = self.client.get(
            base, {"seqid": "chr01", "start": 1, "end": 3, "strand": "-"}
        ).json()
        self.assertEqual(plus["sequence"], "ATG")
        self.assertEqual(minus["sequence"], "CAT")
        self.assertEqual(
            self.client.get(base, {"seqid": "chr01", "start": 0}).status_code, 400
        )
        length = index.summary["sequences"][0]["length"]
        last = self.client.get(
            base, {"seqid": "chr01", "start": length, "end": length}
        ).json()
        self.assertEqual(len(last["sequence"]), 1)
        genes = self.client.get(
            f"/api/omics/assemblies/{index.assembly_id}/genes/",
            {"seqid": "chr01", "start": 303, "end": 303},
        ).json()
        self.assertEqual(genes["total"], 1)
        next_base = self.client.get(
            f"/api/omics/assemblies/{index.assembly_id}/genes/",
            {"seqid": "chr01", "start": 304, "end": 304},
        ).json()
        self.assertEqual(next_base["total"], 0)

    def test_genome_resource_permissions_cover_every_read_surface(self):
        index = self.genome()
        self.client.force_login(self.other)
        self.assertEqual(self.client.get("/api/omics/assemblies/").json()["items"], [])
        for suffix in [
            "summary/",
            "genes/",
            "genes/DEMO-G01001/",
            "sequence/?seqid=chr01",
        ]:
            self.assertEqual(
                self.client.get(
                    f"/api/omics/assemblies/{index.assembly_id}/" + suffix
                ).status_code,
                403,
            )

    def test_genome_bad_parent_fails_without_publishing(self):
        genome = self.root / "bad.fasta"
        genome.write_text(">chr01\nACGTACGT\n")
        gff = self.root / "bad.gff"
        gff.write_text("chr01\tx\texon\t1\t3\t.\t+\t.\tParent=missing\n")
        with self.assertRaises(GenomeError):
            build_index(
                files={"genome": genome, "gff": gff},
                code="BAD",
                name="bad",
                user=self.owner,
                provenance="observed",
            )
        index = GenomeIndex.objects.get(assembly__assembly_code="BAD")
        self.assertEqual(index.status, "failed")
        self.assertEqual(index.assembly.data_resource.status, "inactive")

    def test_browser_genome_import_and_failed_code_are_explicit(self):
        with zipfile.ZipFile(io.BytesIO(genome_example())) as z:
            files = {
                r: SimpleUploadedFile(n, z.read(n))
                for r, n in [
                    ("genome", "genome.fasta"),
                    ("gff", "annotation.gff3"),
                    ("protein", "protein.faa"),
                ]
            }
        response = self.client.post(
            "/api/omics/import/",
            {
                "code": "BROWSER",
                "name": "browser synthetic",
                "provenance": "synthetic",
                **files,
            },
        )
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(response.json()["geneCount"], 60)

    def test_genome_preflight_checks_content_without_creating_business_records(self):
        from apps.omics.models import (
            GenomeAssembly,
            GenomeDataset,
            GenomeAnnotation,
            GenomeSequenceFile,
        )

        models = [
            DataResource,
            GenomeAssembly,
            GenomeDataset,
            GenomeAnnotation,
            GenomeIndex,
            GenomeSequenceFile,
        ]
        before = [m.objects.count() for m in models]

        def payload():
            with zipfile.ZipFile(io.BytesIO(genome_example())) as z:
                return {
                    "code": "PREFLIGHT",
                    "name": "preflight",
                    "provenance": "synthetic",
                    "validateOnly": "true",
                    "genome": SimpleUploadedFile("genome.fa", z.read("genome.fasta")),
                    "gff": SimpleUploadedFile(
                        "annotation.gff3", z.read("annotation.gff3")
                    ),
                }

        response = self.client.post("/api/omics/import/", payload())
        self.assertEqual(response.status_code, 200, response.content)
        self.assertTrue(response.json()["valid"])
        self.assertEqual(response.json()["rowCount"], 60)
        self.assertEqual([m.objects.count() for m in models], before)
        invalid = payload()
        invalid["gff"] = SimpleUploadedFile(
            "wrong.gff3", b"absent\tx\tgene\t1\t3\t.\t+\t.\tID=x\n"
        )
        self.assertEqual(
            self.client.post("/api/omics/import/", invalid).status_code, 400
        )
        self.assertEqual([m.objects.count() for m in models], before)
        final = payload()
        final["validateOnly"] = "false"
        self.assertEqual(self.client.post("/api/omics/import/", final).status_code, 201)

    def test_batch_genome_matching_gzip_validation_and_final_import(self):
        def payload(validate_only):
            with zipfile.ZipFile(io.BytesIO(genome_example())) as z:
                uploads = [
                    SimpleUploadedFile(
                        "GWHDEMO0001." + label + ".gz", gzip.compress(z.read(source))
                    )
                    for label, source in [
                        ("RNA.fasta", "rna.fasta"),
                        ("CDS.fasta", "cds.fasta"),
                        ("gff", "annotation.gff3"),
                        ("genome.fasta", "genome.fasta"),
                        ("Protein.faa", "protein.faa"),
                    ]
                ]
            return {
                "code": "BATCH",
                "name": "Batch",
                "provenance": "synthetic",
                "validateOnly": str(validate_only).lower(),
                "files": uploads,
            }

        before = DataResource.objects.count()
        response = self.client.post("/api/omics/import/", payload(True))
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["rowCount"], 60)
        self.assertEqual(
            {x["role"] for x in response.json()["matchedFiles"]},
            {"genome", "gff", "protein", "rna", "cds"},
        )
        self.assertEqual(DataResource.objects.count(), before)
        response = self.client.post("/api/omics/import/", payload(False))
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(response.json()["geneCount"], 60)

    def test_batch_genome_rejects_duplicate_unknown_mixed_and_invalid_content(self):
        for names, message in [
            (["FG.genome.fa", "FG.gff", "FG.genome.fasta"], "重复"),
            (["unknown.fasta", "FG.gff"], "无法明确识别"),
            (["GWHAAAA0001.genome.fa", "GWHBBBB0001.gff"], "不同GWH"),
            (["FG.genome.fa", "FG.RNA.fasta"], "缺少"),
            (["FG.genome.fa", "FG.gff"], "FASTA"),
        ]:
            response = self.client.post(
                "/api/omics/import/",
                {
                    "code": "BAD-BATCH",
                    "name": "bad",
                    "provenance": "observed",
                    "validateOnly": "true",
                    "files": [SimpleUploadedFile(name, b"invalid") for name in names],
                },
            )
            self.assertEqual(response.status_code, 400, response.content)
            self.assertIn(message, response.json()["detail"])

    def test_sparse_molecular_preflight_preview_and_commit_revalidation(self):
        raw = b"result_id,sample_code,marker,unit,value\nR1,001,DNA,ng/uL,0\nR2,002,DNA,ng/uL,\n"
        args = {
            "name": "sparse",
            "kind": "molecular",
            "provenance": "observed",
            "validateOnly": "true",
        }
        before = DataResource.objects.count()
        response = self.client.post(
            "/api/analysis/import/",
            {**args, "file": SimpleUploadedFile("data.csv.gz", gzip.compress(raw))},
        )
        self.assertEqual(response.status_code, 201, response.content)
        result = response.json()
        self.assertIsNone(result["resourceId"])
        self.assertEqual(result["rowCount"], 2)
        self.assertEqual(
            result["rows"][0][result["columns"].index("sample_code")], "001"
        )
        self.assertEqual(result["rows"][0][result["columns"].index("value")], "0")
        self.assertEqual(DataResource.objects.count(), before)
        bad = raw.replace(b"R2,002", b"R1,002")
        response = self.client.post(
            "/api/analysis/import/",
            {
                **args,
                "validateOnly": "false",
                "file": SimpleUploadedFile("data.csv", bad),
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(DataResource.objects.count(), before)

    def test_excel_sheet_selection_and_optional_traits_plot(self):
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(["unrelated"])
        worksheet = workbook.create_sheet("性状记录")
        worksheet.append(["species", "height_m"])
        worksheet.append(["胡杨", 0])
        stream = io.BytesIO()
        workbook.save(stream)
        args = {
            "name": "sheet",
            "kind": "traits",
            "provenance": "observed",
            "validateOnly": "true",
        }
        response = self.client.post(
            "/api/analysis/import/",
            {
                **args,
                "sheetName": "性状记录",
                "file": SimpleUploadedFile("data.xlsx", stream.getvalue()),
            },
        )
        self.assertEqual(response.status_code, 201, response.content)
        self.assertTrue(response.json()["valid"])
        self.assertEqual(response.json()["rowCount"], 1)
        wrong = self.client.post(
            "/api/analysis/import/",
            {
                **args,
                "sheetName": "不存在",
                "file": SimpleUploadedFile("data.xlsx", stream.getvalue()),
            },
        )
        self.assertEqual(wrong.status_code, 400)

    def test_gzip_input_crc_and_expansion_budget(self):
        genome = self.root / "tiny.fa.gz"
        genome.write_bytes(gzip.compress(b">chr01\nACGTACGT\n"))
        gff = self.root / "tiny.gff"
        gff.write_text("chr01\tx\tgene\t1\t3\t.\t+\t.\tID=gene1\n")
        with self.assertRaises(GenomeError):
            build_index(
                files={"genome": genome, "gff": gff},
                code="BUDGET",
                name="budget",
                user=self.owner,
                provenance="observed",
                max_expanded=4,
            )
